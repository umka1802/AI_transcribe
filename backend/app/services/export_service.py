import io
from docx import Document
from fpdf import FPDF


class ExportService:
    @staticmethod
    def export_txt(text: str) -> bytes:
        return text.encode("utf-8")

    @staticmethod
    def export_docx(text: str) -> bytes:
        doc = Document()
        for paragraph in text.split("\n"):
            doc.add_paragraph(paragraph)
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_pdf(text: str) -> bytes:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "", 12)
        for line in text.split("\n"):
            pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
        buffer = io.BytesIO()
        pdf.output(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_srt(text: str, segments: list = None) -> bytes:
        if not segments:
            lines = text.strip().split("\n")
            srt_lines = []
            for i, line in enumerate(lines, 1):
                srt_lines.append(str(i))
                srt_lines.append("00:00:00,000 --> 00:00:00,000")
                srt_lines.append(line)
                srt_lines.append("")
            return "\n".join(srt_lines).encode("utf-8")

        srt_lines = []
        for i, seg in enumerate(segments, 1):
            start = seg.get("start", 0)
            end = seg.get("end", 0)
            seg_text = seg.get("text", "").strip()
            if not seg_text:
                continue
            start_str = ExportService._format_srt_time(start)
            end_str = ExportService._format_srt_time(end)
            srt_lines.append(str(i))
            srt_lines.append(f"{start_str} --> {end_str}")
            srt_lines.append(seg_text)
            srt_lines.append("")
        return "\n".join(srt_lines).encode("utf-8")

    @staticmethod
    def _format_srt_time(seconds: float) -> str:
        h = int(seconds // 3600)
        m = int((seconds % 3600) // 60)
        s = int(seconds % 60)
        ms = int((seconds % 1) * 1000)
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


export_service = ExportService()
